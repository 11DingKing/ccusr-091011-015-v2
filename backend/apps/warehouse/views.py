"""
仓库管理视图
"""
import logging
import io
from decimal import Decimal, InvalidOperation

from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from apps.core.exceptions import BusinessException
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    StorageZone, ZoneRuleVersion, ZoneRuleIncompatibility, ZoneCapacity,
    CapacityOverride, StockPlacement, ZoneReservation, ZoneTransfer,
    ZoneValidationLog,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    StorageZoneSerializer, StorageZoneCreateSerializer,
    ZoneRuleVersionSerializer, ZoneRuleVersionCreateSerializer,
    ZoneCapacitySerializer, CapacityOverrideSerializer,
    CapacityOverrideCreateSerializer, StockPlacementSerializer,
    ZoneReservationSerializer, ZoneTransferSerializer,
    ZoneValidationLogSerializer,
)
from .services import ZoneRuleEngine

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related(
            'goods', 'operator', 'zone', 'rule_version'
        ).order_by('-stock_in_time')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]
        serializer = StockInSerializer(records, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """
        入库（指定保管区）

        与转移、预约共用同一容量/相容性校验口径；
        可携带 reservation_id 按预约履约入库。
        """
        goods_id = request.data.get('goods')
        zone_id = request.data.get('zone')
        quantity_raw = request.data.get('quantity')
        reservation_id = request.data.get('reservation')

        if not goods_id or not zone_id or quantity_raw is None:
            return error_response(message='货物、保管区、数量均为必填项')

        try:
            quantity = Decimal(str(quantity_raw))
        except (InvalidOperation, ValueError):
            return error_response(message='数量格式不正确')
        if quantity <= 0:
            return error_response(message='入库数量必须大于零')

        goods = Goods.objects.filter(pk=goods_id, is_active=True).first()
        if not goods:
            return error_response(message='货物不存在或已停用')
        zone = StorageZone.objects.filter(pk=zone_id).first()
        if not zone:
            return error_response(message='保管区不存在')

        reservation = None
        if reservation_id:
            reservation = ZoneReservation.objects.filter(
                pk=reservation_id, zone=zone, goods=goods
            ).first()
            if not reservation:
                return error_response(message='预约不存在或与货物、保管区不匹配')
            if reservation.quantity != quantity:
                return error_response(message='入库数量须与预约数量一致')

        try:
            stock_in = ZoneRuleEngine.stock_in(
                zone=zone, goods=goods, quantity=quantity,
                operator=request.user,
                batch_no=request.data.get('batch_no', ''),
                supplier=request.data.get('supplier', ''),
                remark=request.data.get('remark', ''),
                reservation=reservation,
            )
        except BusinessException as e:
            return error_response(message=e.message, code=e.code)

        return success_response(data=StockInSerializer(stock_in).data, message='入库成功')


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 保管区与容量规则 ====================


def _require_admin(request):
    """区域规则与临时超限属于受控配置，仅管理员可操作"""
    if not request.user.is_admin:
        return error_response(message='无权限操作', code=403)
    return None


class ZoneListView(APIView):
    """保管区列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StorageZone.objects.all().order_by('code')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        zones = queryset[start:end]
        serializer = StorageZoneSerializer(zones, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建保管区（限管理员）"""
        denied = _require_admin(request)
        if denied:
            return denied

        serializer = StorageZoneCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))

        zone = StorageZone.objects.create(
            name=serializer.validated_data['name'],
            code=serializer.validated_data['code'],
            description=serializer.validated_data.get('description', ''),
            created_by=request.user,
        )
        ZoneCapacity.objects.create(zone=zone)

        logger.info(f"User {request.user.username} created zone {zone.code}")
        return success_response(data=StorageZoneSerializer(zone).data, message='创建成功')


class ZoneDetailView(APIView):
    """保管区详情视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)
        return success_response(data=StorageZoneSerializer(zone).data)

    def put(self, request, pk):
        """更新保管区基本信息（限管理员；规则调整走版本接口）"""
        denied = _require_admin(request)
        if denied:
            return denied

        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        serializer = StorageZoneCreateSerializer(
            data=request.data, context={'instance': zone}
        )
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))

        zone.name = serializer.validated_data['name']
        zone.code = serializer.validated_data['code']
        zone.description = serializer.validated_data.get('description', '')
        if 'is_active' in request.data:
            zone.is_active = bool(request.data['is_active'])
        zone.save()

        logger.info(f"User {request.user.username} updated zone {zone.code}")
        return success_response(data=StorageZoneSerializer(zone).data, message='更新成功')

    def delete(self, request, pk):
        """删除保管区（限管理员；存在存放或记录时仅可停用）"""
        denied = _require_admin(request)
        if denied:
            return denied

        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        if zone.placements.filter(quantity__gt=0).exists():
            return error_response(message='保管区内仍有物资，无法删除')
        if zone.reservations.filter(status='active').exists():
            return error_response(message='保管区存在生效中的预约，无法删除')
        if zone.validation_logs.exists() or zone.rule_versions.exists():
            return error_response(message='保管区已产生规则或校验记录，仅可停用不可删除')

        code = zone.code
        zone.delete()
        logger.info(f"User {request.user.username} deleted zone {code}")
        return success_response(message='删除成功')


class ZoneRuleListView(APIView):
    """保管区规则版本视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        """规则版本列表（历史版本全量保留）"""
        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        versions = zone.rule_versions.prefetch_related(
            'incompatibilities', 'allowed_categories'
        ).order_by('-version')
        serializer = ZoneRuleVersionSerializer(versions, many=True)
        return success_response(data=serializer.data)

    def post(self, request, pk):
        """新增规则版本（限管理员；旧版本不可变，历史记录按发生时版本解释）"""
        denied = _require_admin(request)
        if denied:
            return denied

        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        serializer = ZoneRuleVersionCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        data = serializer.validated_data
        last_version = zone.rule_versions.order_by('-version').first()
        version = ZoneRuleVersion.objects.create(
            zone=zone,
            version=(last_version.version + 1) if last_version else 1,
            max_weight=data['max_weight'],
            max_items=data['max_items'],
            effective_from=data.get('effective_from') or timezone.now(),
            created_by=request.user,
        )
        version.allowed_categories.set(data.get('allowed_category_ids', []))
        for category_a_id, category_b_id in data.get('incompatible_pairs', []):
            ZoneRuleIncompatibility.objects.create(
                rule_version=version,
                category_a_id=category_a_id,
                category_b_id=category_b_id,
            )

        logger.info(
            f"User {request.user.username} created rule version "
            f"{zone.code} v{version.version}"
        )
        return success_response(
            data=ZoneRuleVersionSerializer(version).data, message='规则版本已创建'
        )


class ZoneOccupancyView(APIView):
    """保管区占用情况视图（当前生效口径下的占用与剩余）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        ZoneRuleEngine.expire_due_reservations(zone)
        capacity, _ = ZoneCapacity.objects.get_or_create(zone=zone)
        version, max_weight, max_items, overrides = (
            ZoneRuleEngine.get_effective_limits(zone)
        )

        data = {
            'zone': StorageZoneSerializer(zone).data,
            'capacity': ZoneCapacitySerializer(capacity).data,
            'rule_version': (
                ZoneRuleVersionSerializer(version).data if version else None
            ),
            'effective_limits': {
                'max_weight': str(max_weight) if max_weight is not None else None,
                'max_items': str(max_items) if max_items is not None else None,
            },
            'remaining': {
                'weight': (
                    str(max_weight - capacity.current_weight - capacity.reserved_weight)
                    if max_weight is not None else None
                ),
                'items': (
                    str(max_items - capacity.current_items - capacity.reserved_items)
                    if max_items is not None else None
                ),
            },
            'active_overrides': CapacityOverrideSerializer(overrides, many=True).data,
        }
        return success_response(data=data)


class ZoneOverrideListView(APIView):
    """临时超限批准视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)
        overrides = zone.overrides.order_by('-created_at')
        return success_response(
            data=CapacityOverrideSerializer(overrides, many=True).data
        )

    def post(self, request, pk):
        """批准临时超限（限管理员，必须设置失效时间）"""
        denied = _require_admin(request)
        if denied:
            return denied

        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        serializer = CapacityOverrideCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        override = CapacityOverride.objects.create(
            zone=zone,
            extra_weight=serializer.validated_data['extra_weight'],
            extra_items=serializer.validated_data['extra_items'],
            reason=serializer.validated_data['reason'],
            expires_at=serializer.validated_data['expires_at'],
            approved_by=request.user,
        )
        logger.info(
            f"User {request.user.username} approved override for zone "
            f"{zone.code} until {override.expires_at}"
        )
        return success_response(
            data=CapacityOverrideSerializer(override).data, message='临时超限已批准'
        )


class ZoneOverrideRevokeView(APIView):
    """撤销临时超限批准（限管理员）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        denied = _require_admin(request)
        if denied:
            return denied

        override = CapacityOverride.objects.filter(pk=pk).first()
        if not override:
            return error_response(message='临时超限批准不存在', code=404)
        if override.is_revoked:
            return error_response(message='该批准已被撤销')

        override.is_revoked = True
        override.save(update_fields=['is_revoked'])
        logger.info(
            f"User {request.user.username} revoked override {override.id}"
        )
        return success_response(
            data=CapacityOverrideSerializer(override).data, message='已撤销'
        )


class ReservationListView(APIView):
    """容量预约视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ZoneReservation.objects.select_related(
            'zone', 'goods', 'created_by', 'rule_version'
        ).order_by('-created_at')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        serializer = ZoneReservationSerializer(queryset[start:end], many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建预约（与入库、转移同一校验口径）"""
        goods_id = request.data.get('goods')
        zone_id = request.data.get('zone')
        quantity_raw = request.data.get('quantity')
        expires_at_raw = request.data.get('expires_at')

        if not goods_id or not zone_id or quantity_raw is None or not expires_at_raw:
            return error_response(message='货物、保管区、数量、失效时间均为必填项')

        try:
            quantity = Decimal(str(quantity_raw))
        except (InvalidOperation, ValueError):
            return error_response(message='数量格式不正确')
        if quantity <= 0:
            return error_response(message='预约数量必须大于零')

        goods = Goods.objects.filter(pk=goods_id, is_active=True).first()
        if not goods:
            return error_response(message='货物不存在或已停用')
        zone = StorageZone.objects.filter(pk=zone_id).first()
        if not zone:
            return error_response(message='保管区不存在')

        expires_at = timezone.datetime.fromisoformat(str(expires_at_raw))
        if timezone.is_naive(expires_at):
            expires_at = timezone.make_aware(expires_at)

        try:
            reservation = ZoneRuleEngine.reserve(
                zone=zone, goods=goods, quantity=quantity,
                expires_at=expires_at, operator=request.user,
            )
        except BusinessException as e:
            return error_response(message=e.message, code=e.code)

        return success_response(
            data=ZoneReservationSerializer(reservation).data, message='预约成功'
        )


class ReservationCancelView(APIView):
    """取消预约视图（释放预约占用）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        reservation = ZoneReservation.objects.filter(pk=pk).first()
        if not reservation:
            return error_response(message='预约不存在', code=404)

        try:
            ZoneRuleEngine.cancel_reservation(reservation, request.user)
        except BusinessException as e:
            return error_response(message=e.message, code=e.code)

        return success_response(
            data=ZoneReservationSerializer(reservation).data, message='预约已取消'
        )


class TransferListView(APIView):
    """区域转移视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ZoneTransfer.objects.select_related(
            'goods', 'from_zone', 'to_zone', 'operator', 'rule_version'
        ).order_by('-created_at')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        serializer = ZoneTransferSerializer(queryset[start:end], many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """转移（与入库、预约同一校验口径，源区释放与目标区占用同事务）"""
        goods_id = request.data.get('goods')
        from_zone_id = request.data.get('from_zone')
        to_zone_id = request.data.get('to_zone')
        quantity_raw = request.data.get('quantity')

        if not goods_id or not from_zone_id or not to_zone_id or quantity_raw is None:
            return error_response(message='货物、转出区、转入区、数量均为必填项')

        try:
            quantity = Decimal(str(quantity_raw))
        except (InvalidOperation, ValueError):
            return error_response(message='数量格式不正确')
        if quantity <= 0:
            return error_response(message='转移数量必须大于零')

        goods = Goods.objects.filter(pk=goods_id, is_active=True).first()
        if not goods:
            return error_response(message='货物不存在或已停用')
        from_zone = StorageZone.objects.filter(pk=from_zone_id).first()
        to_zone = StorageZone.objects.filter(pk=to_zone_id).first()
        if not from_zone or not to_zone:
            return error_response(message='保管区不存在')

        try:
            transfer = ZoneRuleEngine.transfer(
                goods=goods, from_zone=from_zone, to_zone=to_zone,
                quantity=quantity, operator=request.user,
            )
        except BusinessException as e:
            return error_response(message=e.message, code=e.code)

        return success_response(
            data=ZoneTransferSerializer(transfer).data, message='转移成功'
        )


class ZoneValidationLogListView(APIView):
    """容量校验记录视图（历史记录按发生时规则版本解释）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        queryset = zone.validation_logs.select_related(
            'rule_version', 'goods', 'category', 'actor', 'override'
        ).order_by('-created_at')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        serializer = ZoneValidationLogSerializer(queryset[start:end], many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class ZonePlacementListView(APIView):
    """保管区存放明细视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        zone = StorageZone.objects.filter(pk=pk).first()
        if not zone:
            return error_response(message='保管区不存在', code=404)

        placements = zone.placements.filter(quantity__gt=0).select_related('goods')
        serializer = StockPlacementSerializer(placements, many=True)
        return success_response(data=serializer.data)
